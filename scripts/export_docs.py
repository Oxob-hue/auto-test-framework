"""文档导出脚本：把 docs 下的 Markdown 导出为 Excel/CSV 与可打印 HTML（供打印为 PDF）。

用法：
    python scripts/export_docs.py

产出（docs/export/ 目录）：
    测试用例表.xlsx            # 多个 Sheet：接口用例 / Web 用例 / 统计与执行
    测试用例表.csv             # UTF-8 BOM，Excel 直接双击不乱码
    html/测试计划.html
    html/测试用例表.html
    html/测试报告.html          # 用浏览器打开后 Ctrl+P 即可另存为 PDF

设计说明：解析 Markdown 表格（含中文表头）→ 结构化数据 → 一份数据两处复用
（Excel 供人查看、HTML 供打印），避免文档与表格两份内容不一致。
"""
import csv
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import markdown
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
DOCS = ROOT / "docs"
OUT_DIR = DOCS / "export"
HTML_DIR = OUT_DIR / "html"

PRINT_DOCS = ["测试计划.md", "测试用例表.md", "测试报告.md"]

PRINT_CSS = """
@page { size: A4; margin: 14mm 12mm; }
body { font-family: "Microsoft YaHei", "PingFang SC", sans-serif; color:#222;
       font-size: 11.5px; line-height: 1.6; margin: 0 auto; max-width: 190mm; }
h1 { font-size: 19px; border-bottom: 2px solid #2f6fed; padding-bottom: 6px; }
h2 { font-size: 14px; color:#1a4fbf; margin-top: 14px; border-left: 4px solid #2f6fed; padding-left: 7px; }
h3 { font-size: 12.5px; margin-top: 10px; }
table { border-collapse: collapse; width: 100%; margin: 6px 0 10px; }
th, td { border: 1px solid #bbb; padding: 3px 5px; font-size: 10.5px; vertical-align: top; }
th { background: #eef3ff; font-weight: 600; }
tr { page-break-inside: avoid; }
code { background: #f5f5f5; padding: 1px 3px; border-radius: 3px; font-size: 10.5px; }
blockquote { border-left: 3px solid #ccc; color: #555; margin: 6px 0; padding: 2px 10px; }
pre { background:#f7f7f7; padding:6px 8px; overflow:auto; }
hr { border: none; border-top: 1px dashed #ccc; margin: 12px 0; }
ul, ol { margin: 4px 0 4px 20px; padding: 0; }
"""


def md_to_html(md_path: Path) -> Path:
    """Markdown → 带打印样式的独立 HTML 文件。"""
    text = md_path.read_text(encoding="utf-8")
    body = markdown.markdown(text, extensions=["tables", "fenced_code", "sane_lists"])
    html_text = (
        "<!DOCTYPE html>\n<html lang=\"zh-CN\"><head><meta charset=\"utf-8\">"
        f"<title>{md_path.stem}</title><style>{PRINT_CSS}</style></head>"
        f"<body>{body}</body></html>"
    )
    HTML_DIR.mkdir(parents=True, exist_ok=True)
    out = HTML_DIR / f"{md_path.stem}.html"
    out.write_text(html_text, encoding="utf-8")
    return out


def parse_tables(md_path: Path):
    """解析 Markdown 中的表格，返回 [(小标题, 表头, 数据行...), ...]。"""
    lines = md_path.read_text(encoding="utf-8").splitlines()
    tables, heading = [], ""
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if line.startswith("#"):
            heading = line.lstrip("#").strip()
        if line.startswith("|") and i + 1 < len(lines) and re.fullmatch(r"\|[\s:|-]+\|", lines[i + 1].strip()):
            header = [c.strip() for c in line.strip("|").split("|")]
            rows, j = [], i + 2
            while j < len(lines) and lines[j].strip().startswith("|"):
                rows.append([c.strip().replace("<br>", "\n") for c in lines[j].strip().strip("|").split("|")])
                j += 1
            tables.append((heading, header, rows))
            i = j
            continue
        i += 1
    return tables


def safe_sheet_name(name: str, used: set) -> str:
    title = re.sub(r"[\\/*?:\[\]]", "", name) or "Sheet"
    title = re.sub(r"^\s*[一二三四五六七八九十]+、\s*", "", title)[:28] or "Sheet"
    base, idx = title, 1
    while title in used:
        idx += 1
        title = f"{base}_{idx}"[:31]
    used.add(title)
    return title


def write_xlsx(tables, path: Path) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    used: set = set()
    thin = Side(style="thin", color="BBBBBB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor="EEF3FF")
    head_font = Font(bold=True, size=10)
    body_font = Font(size=10)

    for heading, header, rows in tables:
        ws = wb.create_sheet(safe_sheet_name(heading, used))
        ws.append(header)
        for r in rows:
            ws.append(r)
        # 表头样式
        for col in range(1, len(header) + 1):
            cell = ws.cell(row=1, column=col)
            cell.font, cell.fill, cell.border = head_font, head_fill, border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        # 正文样式 + 列宽
        for col in range(1, len(header) + 1):
            max_len = len(str(header[col - 1]))
            for row in range(2, ws.max_row + 1):
                cell = ws.cell(row=row, column=col)
                cell.font, cell.border = body_font, border
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                max_len = max(max_len, max((len(s) for s in str(cell.value or "").splitlines()), default=0))
            ws.column_dimensions[get_column_letter(col)].width = min(max(10, max_len * 1.6), 46)
        ws.freeze_panes = "A2"
        ws.row_dimensions[1].height = 30
    wb.save(path)


def write_csv(tables, path: Path) -> None:
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        for heading, header, rows in tables:
            writer.writerow([f"【{heading}】"])
            writer.writerow(header)
            writer.writerows(rows)
            writer.writerow([])


def main() -> int:
    OUT_DIR.mkdir(parents=True, exist_ok=True)

    # 1) 打印用 HTML（三份文档）
    for name in PRINT_DOCS:
        src = DOCS / name
        if src.exists():
            print("HTML:", md_to_html(src).relative_to(ROOT))

    # 2) 用例表 → Excel + CSV
    case_doc = DOCS / "测试用例表.md"
    tables = parse_tables(case_doc)
    xlsx, csv_path = OUT_DIR / "测试用例表.xlsx", OUT_DIR / "测试用例表.csv"
    write_xlsx(tables, xlsx)
    write_csv(tables, csv_path)
    print(f"XLSX: {xlsx.relative_to(ROOT)} （{len(tables)} 个 Sheet，共 {sum(len(r) for _,_,r in tables)} 行）")
    print(f"CSV : {csv_path.relative_to(ROOT)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
